"""Real Office files and the bundled SDK connection through normal MCP policy."""
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest

from hanni.dashboard.features import DesktopFeatures
from hanni.dashboard.mcp_hub import MCPHub
from hanni.mcp.office import OfficeFiles, READ_TOOLS, office_root
from hanni.tools.base import ConfirmationRequired, ToolPermissionError, ToolRequest
from tests.unit.test_mcp_hub_api import body, config_for, http_call, server


def test_word_and_slides_roundtrip_and_stale_edits(tmp_path):
    files = OfficeFiles(tmp_path)
    for name, create, append in [
        ('report.docx', lambda: files.word_create('report.docx', ['Quarterly update']),
         lambda digest: files.word_append('report.docx', ['Completed the rollout'], digest)),
        ('review.pptx', lambda: files.powerpoint_create('review.pptx', [{'title': 'Quarterly update', 'body': 'On track'}]),
         lambda digest: files.powerpoint_append('review.pptx', 'Completed the rollout', 'Next steps', digest)),
    ]:
        created = create()
        assert created['saved'] and 'Quarterly update' in files.read_document(name)['text']
        with pytest.raises(FileExistsError): create()
        updated = append(created['sha256'])
        assert updated['sha256'] != created['sha256']
        assert 'Completed the rollout' in files.read_document(name)['text']
        before = (tmp_path / name).read_bytes()
        with pytest.raises(ValueError, match='changed'): append(created['sha256'])
        assert (tmp_path / name).read_bytes() == before


def test_excel_updates_only_requested_cells_and_keeps_formulas_literal(tmp_path):
    from openpyxl import load_workbook
    files = OfficeFiles(tmp_path)
    created = files.excel_create('budget.xlsx', [['Team', 'Budget', 'Comment'], ['Sales', 120, '=1+1']], 'Budget')
    book = load_workbook(tmp_path / 'budget.xlsx')
    assert book['Budget']['C2'].data_type == 's'
    book.close()
    files.excel_write('budget.xlsx', 'Budget', 'B2', [[150, None]], created['sha256'])
    book = load_workbook(tmp_path / 'budget.xlsx')
    assert list(book['Budget'].values) == [('Team', 'Budget', 'Comment'), ('Sales', 150, None)]
    book.close()
    digest = files.read_document('budget.xlsx')['sha256']
    files.excel_write('budget.xlsx', 'Budget', 'C2', [['=SUM(B2:B3)']], digest, allow_formulas=True)
    book = load_workbook(tmp_path / 'budget.xlsx')
    assert book['Budget']['C2'].data_type == 'f'
    book.close()
    with pytest.raises(ValueError, match='10,000 rows'):
        files.excel_write('budget.xlsx', 'Budget', 'A10001', [[1]], files.read_document('budget.xlsx')['sha256'])


def test_folder_scope_and_pagination(tmp_path):
    folder = tmp_path / 'documents'; folder.mkdir()
    files = OfficeFiles(folder)
    (folder / 'report.csv').write_text('column\n' + 'data\n' * 20)
    (folder / 'unlisted.exe').write_text('not a document')
    assert [item['path'] for item in files.list_documents()['files']] == ['report.csv']
    read = files.read_document('report.csv', limit=10)
    assert read['next_offset'] == 10
    assert files.read_document('report.csv', offset=10, limit=10)['text'] == (folder / 'report.csv').read_text()[10:20]
    for path in ['../outside.docx', str(tmp_path / 'outside.docx'), '.ssh/id_rsa']:
        with pytest.raises(ValueError): files.path(path)
    for path in [str(Path(tmp_path.anchor)), str(Path.home()), 'relative/folder']:
        with pytest.raises(ValueError): office_root(path)
    with pytest.raises(ValueError): files.read_document('unlisted.exe')
    with pytest.raises(ValueError): files.word_create('wrong.txt', ['text'])


def test_changed_document_during_save_leaves_original_and_no_temp_files(tmp_path, monkeypatch):
    from docx import Document
    files = OfficeFiles(tmp_path)
    created = files.word_create('report.docx', ['original'])
    document = Document(BytesIO((tmp_path / 'report.docx').read_bytes()))
    actual_read = files.read_bytes
    monkeypatch.setattr(files, 'read_bytes', lambda target: b'changed')
    with pytest.raises(ValueError, match='No changes were saved'):
        files.save(tmp_path / 'report.docx', document, expected=created['sha256'])
    assert not list(tmp_path.glob('*.tmp'))
    monkeypatch.setattr(files, 'read_bytes', actual_read)
    assert files.read_document('report.docx')['sha256'] == created['sha256']


def test_edits_preflight_xml_parts_even_with_unusual_extensions(tmp_path):
    import hashlib
    from zipfile import ZipFile
    from defusedxml.common import DefusedXmlException
    files = OfficeFiles(tmp_path)
    files.word_create('report.docx', ['original'])
    source = (tmp_path / 'report.docx').read_bytes()
    out = BytesIO()
    with ZipFile(BytesIO(source)) as before, ZipFile(out, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry.filename)
            if entry.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', b'<Override PartName="/word/extra.bin" ContentType="application/test+xml"/></Types>')
            after.writestr(entry, data)
        after.writestr('word/extra.bin', '<!DOCTYPE r [<!ENTITY s "expanded">]><r>&s;</r>')
    raw = out.getvalue()
    (tmp_path / 'report.docx').write_bytes(raw)
    with pytest.raises(DefusedXmlException):
        files.word_append('report.docx', ['text'], hashlib.sha256(raw).hexdigest())
    assert (tmp_path / 'report.docx').read_bytes() == raw


@pytest.fixture
def hub(tmp_path):
    config = config_for(tmp_path / 'data')
    instance = MCPHub(config, DesktopFeatures(config))
    instance.manager.attach_registry(instance.features.registry)
    yield instance
    instance.close()


def test_setup_discovers_real_tools_and_preserves_execution_approval(hub, tmp_path):
    folder = tmp_path / 'documents'
    settings = {'folder': str(folder), 'enable_read_tools': True, 'auto_connect': False}
    request = body('office_setup', settings)
    created = hub.mutate(request)
    assert created['connected'] and folder.is_dir()
    assert hub.mutate(request) == created
    assert hub.mutate(body('office_setup', settings))['server_id'] == created['server_id']
    assert len(hub.snapshot()['servers']) == 1
    server = hub.snapshot()['servers'][0]
    assert server['discovered_count'] == 8 and server['enabled_count'] == 2
    assert {tool['original_name'] for tool in server['tools'] if tool['enabled']} == READ_TOOLS
    assert all(tool['valid'] for tool in server['tools'])
    read = next(tool for tool in server['tools'] if tool['original_name'] == 'office_list_documents')
    registry = hub.features.registry
    assert registry.execute(ToolRequest(read['internal_id'], {})).data['structured_content']['files'] == []
    write = next(tool for tool in server['tools'] if tool['original_name'] == 'word_create_document')
    hub.mutate(body('approve_tools', {'server_id': server['id'], 'tools': [{
        'original_name': write['original_name'], 'definition_hash': write['definition_hash'], 'classification': 'write'}]}))
    arguments = {'path': 'minutes.docx', 'paragraphs': ['Meeting notes']}
    with pytest.raises(ConfirmationRequired): registry.execute(ToolRequest(write['internal_id'], arguments))
    assert not (folder / 'minutes.docx').exists()
    result = registry.execute(ToolRequest(write['internal_id'], arguments, confirmed=True))
    assert result.ok and (folder / 'minutes.docx').exists()
    duplicate = registry.execute(ToolRequest(write['internal_id'], arguments, confirmed=True))
    assert not duplicate.ok and duplicate.data['is_error']
    # Disconnect + retry reconnects the same server rather than duplicating it.
    hub.manager.disconnect(server['id'])
    assert hub.mutate(body('office_setup', settings))['server_id'] == server['id']
    with pytest.raises(ValueError, match='different settings'):
        hub.mutate(body('office_setup', {**settings, 'folder': str(tmp_path / 'other')}))
    assert not (tmp_path / 'other').exists()


def test_setup_can_leave_every_tool_disabled_and_pin_reconnect(hub, tmp_path):
    settings = {'folder': str(tmp_path / 'documents'), 'enable_read_tools': False, 'auto_connect': True}
    created = hub.mutate(body('office_setup', settings))
    server = hub.snapshot()['servers'][0]
    assert server['enabled_count'] == 0 and server['config']['auto_connect']
    assert created['server_id'] == server['id']
    assert hub.office_status()['default_folder'] == settings['folder']
    assert hub.office_status()['auto_connect'] and not hub.office_status()['enable_read_tools']


def test_repeated_local_setup_applies_the_reading_checkbox(hub, tmp_path):
    settings = {'folder': str(tmp_path / 'documents'), 'enable_read_tools': True, 'auto_connect': False}
    hub.mutate(body('office_setup', settings))
    assert hub.snapshot()['servers'][0]['enabled_count'] == 2
    hub.mutate(body('office_setup', {**settings, 'enable_read_tools': False}))
    assert hub.snapshot()['servers'][0]['enabled_count'] == 0


def test_missing_dependencies_do_not_create_a_folder_or_save_server(hub, tmp_path, monkeypatch):
    from hanni.dashboard import office_setup
    actual = office_setup.importlib.util.find_spec
    monkeypatch.setattr(office_setup.importlib.util, 'find_spec', lambda name: None if name == 'pptx' else actual(name))
    assert hub.office_status()['missing_packages'] == ['pptx']
    with pytest.raises(ValueError, match='SETUP_HANNI.cmd'):
        hub.mutate(body('office_setup', {'folder': str(tmp_path / 'docs'), 'enable_read_tools': True, 'auto_connect': False}))
    assert not (tmp_path / 'docs').exists() and not hub.snapshot()['servers']


def test_retry_after_connection_failure_reuses_saved_configuration(hub, tmp_path, monkeypatch):
    settings = {'folder': str(tmp_path / 'docs'), 'enable_read_tools': True, 'auto_connect': False}
    connect = hub.manager.connect
    monkeypatch.setattr(hub.manager, 'connect', lambda *args, **kwargs: hub.manager.snapshot())
    request = body('office_setup', settings)
    with pytest.raises(ValueError, match='settings were saved'):
        hub.mutate(request)
    assert len(hub.snapshot()['servers']) == 1
    monkeypatch.setattr(hub.manager, 'connect', connect)
    assert hub.mutate(request)['connected']
    assert len(hub.snapshot()['servers']) == 1


def test_office_setup_policy_denies_before_creating_folder(tmp_path):
    config = config_for(tmp_path / 'data')
    owner = replace(config.permissions.profiles['owner'], chat_only=True)
    config = replace(config, permissions=replace(config.permissions, profiles={**config.permissions.profiles, 'owner': owner}))
    hub = MCPHub(config, DesktopFeatures(config))
    try:
        with pytest.raises(ToolPermissionError):
            hub.mutate(body('office_setup', {'folder': str(tmp_path / 'docs'), 'enable_read_tools': True, 'auto_connect': False}))
        assert not (tmp_path / 'docs').exists() and not hub.snapshot()['servers']
    finally: hub.close()


def test_office_setup_status_is_authenticated_and_setup_uses_fixed_operation(server, tmp_path):
    assert http_call(server, '/api/mcp/office', authorized=False)[0] == 403
    assert http_call(server, '/api/mcp/office', origin='https://foreign.example')[0] == 403
    code, ready = http_call(server, '/api/mcp/office')
    assert code == 200 and ready['available']
    settings = {'folder': str(tmp_path / 'docs'), 'enable_read_tools': True, 'auto_connect': False}
    request = body('office_setup', settings)
    assert http_call(server, '/api/mcp', method='POST', data=request, authorized=False)[0] == 403
    assert not (tmp_path / 'docs').exists()
    assert http_call(server, '/api/mcp', method='POST', data={**request, 'data': {**settings, 'command': 'arbitrary'}})[0] == 400
    code, result = http_call(server, '/api/mcp', method='POST', data=request)
    assert code == 200 and result['connected']


def test_cloud_setup_requires_installed_workiq_and_preserves_local_connection(hub, tmp_path, monkeypatch):
    local = hub.mutate(body('office_setup', {'folder': str(tmp_path / 'documents'), 'enable_read_tools': True, 'auto_connect': False}))
    for command in ['', 'workiq', str(tmp_path / 'missing' / 'workiq.exe'), str(tmp_path / 'python.exe')]:
        with pytest.raises(ValueError):
            hub.mutate(body('office_cloud_setup', {'command': command, 'auto_connect': False}))
    executable = tmp_path / 'workiq.cmd'
    executable.write_text('installed fixture')
    # Only the process boundary is substituted; fixed setup, persistence, and
    # normal management policy still execute against the real MCP manager.
    def discovery(server_id, **kwargs):
        with hub.manager.lock:
            saved = hub.manager.servers[server_id]
            saved.update(state='connected', enabled=True)
        return {'servers': [hub.manager._public(saved)]}
    monkeypatch.setattr(hub.manager, 'connect', discovery)
    settings = {'command': str(executable), 'auto_connect': False}
    created = hub.mutate(body('office_cloud_setup', settings))
    assert created['connected'] and not created['account_verified']
    saved = hub.manager.servers[created['server_id']]
    assert saved['config']['args'] == ['mcp'] and saved['config']['command'] == str(executable.resolve())
    assert not saved['tools']
    assert len(hub.snapshot()['servers']) == 2
    assert any(item['id'] == local['server_id'] for item in hub.snapshot()['servers'])
    assert hub.mutate(body('office_cloud_setup', settings))['server_id'] == created['server_id']
    with pytest.raises(ValueError, match='different settings'):
        hub.mutate(body('office_cloud_setup', {**settings, 'auto_connect': True}))


def test_cloud_setup_is_authenticated_and_never_accepts_generic_arguments(server, tmp_path):
    request = body('office_cloud_setup', {'command': str(tmp_path / 'workiq.exe'), 'auto_connect': False})
    assert http_call(server, '/api/mcp', method='POST', data=request, authorized=False)[0] == 403
    request['data']['args'] = ['accept-eula']
    assert http_call(server, '/api/mcp', method='POST', data=request)[0] == 400
    assert not server.mcp.snapshot()['servers']

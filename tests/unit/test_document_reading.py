"""Real document fixtures, extraction limits, workspace pagination and routing."""
from datetime import datetime
import base64
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from uuid import uuid4

from openpyxl import Workbook
import pytest

from hanni.dashboard.attachments import extract_text
from hanni.native.local_tools import execute_local, local_tool_specs
from hanni.native.skills import load_skills, match_skills
from hanni.tools.documents import MAX_TEXT
from tests.unit.test_dashboard_chat import chat_server, send


def archive_bytes(parts):
    result = BytesIO()
    with ZipFile(result, 'w', ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return result.getvalue()


def workbook_bytes():
    book = Workbook()
    sheet = book.active; sheet.title = 'Sales Ω'
    sheet.append(['Code', 'Amount', 'Calculated', 'Date'])
    sheet.append(['0012', 42.5, '=B2*2', datetime(2026, 10, 24)])
    sheet['E2'] = False
    sheet['F2'] = '#DIV/0!'
    other = book.create_sheet('Other'); other['A1'] = 'Second sheet'; other.sheet_state = 'hidden'
    result = BytesIO(); book.save(result); book.close()
    # Simulate a value cached by an actual spreadsheet app and incorrect
    # producer dimension metadata; extraction must use actual cell references.
    with ZipFile(BytesIO(result.getvalue())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts['xl/worksheets/sheet1.xml'] = parts['xl/worksheets/sheet1.xml'].replace(
        b'<f>B2*2</f><v />', b'<f>B2*2</f><v>85</v>').replace(
        b'<f>B2*2</f><v></v>', b'<f>B2*2</f><v>85</v>').replace(b'ref="A1:F2"', b'ref="A1:A1"')
    return archive_bytes(parts)


@pytest.mark.parametrize('suffix', ['xlsx', 'xlsm'])
def test_excel_keeps_cells_types_dates_formulas_cache_and_all_sheets(suffix):
    text, media, truncated = extract_text('revenue.' + suffix, workbook_bytes())
    assert not truncated
    assert '[Sheet "Sales Ω"; visible]' in text and '[Sheet "Other"; hidden]' in text
    row = json.loads(next(line for line in text.splitlines() if line.startswith('{"A2"')))
    assert row == {'A2': '0012', 'B2': 42.5, 'C2': {'formula': '=B2*2', 'cached_value': 85},
                   'D2': '2026-10-24T00:00:00', 'E2': False, 'F2': '#DIV/0!'}
    assert 'spreadsheet' in media or 'macroEnabled' in media


def test_xls_reads_real_biff_workbook_without_a_writer_dependency():
    # Small fixture generated once with xlwt 1.3: headers Code/Amount,
    # one row containing the string identifier 0012 and numeric 42.5.
    data = base64.b64decode((Path(__file__).parents[1] / 'fixtures/documents/legacy.xls.b64.txt').read_text())
    text, media, truncated = extract_text('legacy.xls', data)
    assert '"A2": "0012"' in text and '"B2": 42.5' in text
    assert media == 'application/vnd.ms-excel' and not truncated
    assert 'formula expressions are unavailable' in text


OD_START = ('<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
            'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"><office:body>')
OD_END = '</office:body></office:document-content>'


def test_ods_retains_repeated_cell_positions_formulas_and_typed_values():
    raw = archive_bytes({'content.xml': OD_START + '''<office:spreadsheet><table:table table:name="Budget">
      <table:table-row table:number-rows-repeated="2">
        <table:table-cell office:value-type="string"><text:p>0012</text:p></table:table-cell>
        <table:table-cell table:number-columns-repeated="3" office:value-type="float" office:value="42.5" table:formula="of:=21.25*2"><text:p>42.5</text:p></table:table-cell>
      </table:table-row>
      <table:table-row><table:table-cell office:value-type="boolean" office:boolean-value="false"/></table:table-row>
    </table:table></office:spreadsheet>''' + OD_END})
    text, media, truncated = extract_text('budget.ods', raw)
    rows = [json.loads(line) for line in text.splitlines() if line.startswith('{')]
    assert rows[0]['repeat'] == 2 and rows[0]['row'] == 1
    assert rows[0]['cells'][1] == {'column': 2, 'repeat': 3, 'text': '42.5', 'value-type': 'float',
                                   'value': '42.5', 'formula': 'of:=21.25*2'}
    assert rows[1]['row'] == 3 and rows[1]['cells'][0]['boolean-value'] == 'false'
    assert not truncated and media.endswith('spreadsheet')


def test_odt_retains_mixed_text_spaces_tabs_and_breaks():
    raw = archive_bytes({'content.xml': OD_START + '<office:text><text:h>Title</text:h><text:p>Hello <text:span>team</text:span> Ω<text:s text:c="2"/>today<text:tab/>42<text:line-break/>Next</text:p></office:text>' + OD_END})
    assert extract_text('letter.odt', raw)[0] == 'Title\nHello team Ω  today\t42\nNext\n'


def presentation_bytes():
    return archive_bytes({
        'ppt/presentation.xml': '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="r2"/><p:sldId id="257" r:id="r1"/></p:sldIdLst></p:presentation>',
        'ppt/_rels/presentation.xml.rels': '<Relationships><Relationship Id="r1" Target="slides/slide1.xml"/><Relationship Id="r2" Target="slides/slide2.xml"/></Relationships>',
        'ppt/slides/slide1.xml': '<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:r><a:t>Last slide</a:t></a:r></a:p>',
        'ppt/slides/slide2.xml': '<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:r><a:t>First sli</a:t></a:r><a:r><a:t>de Ω</a:t></a:r></a:p>',
    })


def test_powerpoint_follows_presentation_order_not_archive_or_filename_order():
    text, media, truncated = extract_text('slides.pptx', presentation_bytes())
    assert text.index('[Slide 1]\nFirst slide Ω') < text.index('[Slide 2]\nLast slide')
    assert media.endswith('presentation') and not truncated


def test_html_decodes_visible_text_tables_and_entities_without_script_content():
    text, media, truncated = extract_text('export.HTM', b'<html><head><title>Report</title><style>secret-style</style></head><body><h1>Sales &amp; costs</h1><script>do_not_run()</script><table><tr><td>Revenue</td><td>42</td></tr></table><template>hidden-template</template></body></html>')
    assert 'Report' in text and 'Sales & costs' in text and '\tRevenue\t42' in text
    assert all(value not in text for value in ['secret-style', 'do_not_run', 'hidden-template', '<td>'])
    assert media == 'text/html' and not truncated


@pytest.mark.parametrize('name,raw,media', [
    ('rows.csv', 'Code,Note\r\n0012,"a,b\nsecond line"\r\n'.encode(), 'text/csv'),
    ('rows.tsv', b'Code\tAmount\n0012\t42.5', 'text/tab-separated-values'),
    ('data.json', '{"code":"0012","missing":null,"flag":false,"name":"Ω"}'.encode(), 'application/json'),
    ('data.jsonl', b'{"id":1}\n{"id":2}', 'application/x-ndjson'),
])
def test_text_data_preserves_source_structure_and_types(name, raw, media):
    text, actual, truncated = extract_text(name, raw)
    assert text == raw.decode().replace('\r\n', '\n') and actual == media and not truncated


@pytest.mark.parametrize('suffix', ['xlsx', 'xlsm', 'xls', 'ods', 'odt', 'pptx'])
def test_malformed_office_inputs_return_useful_errors(suffix):
    with pytest.raises(ValueError):
        extract_text('broken.' + suffix, b'not an office document')


@pytest.mark.parametrize('suffix,part', [('docx', 'word/document.xml'), ('odt', 'content.xml'),
                                          ('ods', 'content.xml'), ('xlsx', 'xl/workbook.xml')])
def test_office_xml_entities_and_utf16_dtds_are_rejected(suffix, part):
    raw = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE root [<!ENTITY read "untrusted">]><root>&read;</root>'.encode('utf-16')
    with pytest.raises(ValueError):
        extract_text('entities.' + suffix, archive_bytes({part: raw}))


def test_zip_inflation_and_wide_sparse_workbooks_report_limits():
    with pytest.raises(ValueError, match='expanded'):
        extract_text('bomb.ods', archive_bytes({'content.xml': b'x' * (8 * 1024 * 1024 + 1)}))
    book = Workbook(); sheet = book.active
    sheet['A1'] = 'Visible'; sheet['ZZ1'] = 'Outside width'; sheet['A10002'] = 'Outside rows'
    data = BytesIO(); book.save(data); book.close()
    text, _, truncated = extract_text('sparse.xlsx', data.getvalue())
    assert truncated and 'Visible' in text and 'Outside width' not in text and 'Outside rows' not in text


def test_character_cap_is_reported_on_long_open_documents():
    raw = archive_bytes({'content.xml': OD_START + '<office:text><text:p>' + 'x' * (MAX_TEXT + 10) + '</text:p></office:text>' + OD_END})
    text, _, truncated = extract_text('long.odt', raw)
    assert len(text) == MAX_TEXT and truncated


def test_workspace_reader_pages_actual_documents_and_cannot_escape_root(tmp_path):
    path = tmp_path / 'budget.xlsx'; path.write_bytes(workbook_bytes())
    first = execute_local('local_read_document', {'path': 'budget.xlsx', 'limit': 50}, tmp_path)
    assert first['ok'] and first['effects'] == 'none'
    data = first['data']; assert data['offset'] == 0 and data['next_offset'] == 50 and not data['extraction_truncated']
    second = execute_local('local_read_document', {'path': 'budget.xlsx', 'offset': 50, 'limit': 16000}, tmp_path)
    assert second['ok'] and second['data']['next_offset'] is None
    assert data['content'] + second['data']['content'] == extract_text(path.name, path.read_bytes())[0]
    assert not execute_local('local_read_document', {'path': '../outside.pdf'}, tmp_path)['ok']
    assert not execute_local('local_read_document', {'path': 'budget.xlsx', 'offset': -1}, tmp_path)['ok']
    assert not execute_local('local_read_document', {'path': 'budget.xlsx', 'limit': 0}, tmp_path)['ok']
    spec = next(spec for spec in local_tool_specs() if spec['name'] == 'local_read_document')
    from jsonschema import Draft202012Validator
    Draft202012Validator.check_schema(spec['schema']['function']['parameters'])


@pytest.mark.parametrize('prompt,expected', [
    ('What are the exclusions in insurance.pdf?', 'pdf-reading'),
    ('Summarize the attached product.htm export', 'html-reading'),
    ('Find duplicate IDs in orders.jsonl', 'structured-data-reading'),
    ('Review the decision in board.pptx', 'office-document-reading'),
    ('Compare sales across sheets in budget.xlsx', 'spreadsheet-analysis'),
    ('Find errors in the spreadsheet', 'spreadsheet-analysis'),
])
def test_realistic_document_requests_load_the_relevant_bundled_skill(prompt, expected):
    skills = load_skills(metadata_only=True)
    assert '### skill: ' + expected in match_skills(skills, prompt)


@pytest.mark.parametrize('name,source,needle', [
    ('budget.xlsx', workbook_bytes, '"cached_value": 85'),
    ('board.pptx', presentation_bytes, 'First slide Ω'),
    ('report.html', lambda: b'<h1>Quarterly report</h1><script>EXCLUDE_THIS</script>', 'Quarterly report'),
])
def test_new_formats_flow_through_upload_and_chat_context(chat_server, name, source, needle):
    chat_server.chat.native_mode = True
    raw = source()
    status, file = send(chat_server, path='/api/chat/files', body={
        'operation': 'add', 'id': str(uuid4()), 'conversation': 'default', 'name': name,
        'data': base64.b64encode(raw).decode()})
    assert status == 201 and not file['truncated']
    turn = str(uuid4())
    assert send(chat_server, body={'id': turn, 'message': 'Read this', 'attachments': [file['id']]})[0] == 202
    context = chat_server.chat.files.context(turn)
    assert needle in context[0]['text'] and not context[0]['excerpted']
    assert 'EXCLUDE_THIS' not in context[0]['text']
    assert chat_server.chat.files.read(file['id'], 'default') == (name, raw)

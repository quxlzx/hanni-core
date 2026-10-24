"""Bounded, read-only document extraction shared by uploads and workspace tools."""
from __future__ import annotations

from io import BytesIO
from html.parser import HTMLParser
import json
from pathlib import PurePath
from defusedxml import ElementTree as ET
from defusedxml.common import DefusedXmlException
from zipfile import BadZipFile, ZipFile

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_FILES = 5
MAX_TEXT = 200_000
TEXT_EXTENSIONS = frozenset(('.txt .md .markdown .csv .tsv .json .jsonl .xml .yaml .yml '
    '.log .py .js .jsx .ts .tsx .html .htm .css .sql .sh .ps1 .toml .ini .cfg .rst .tex '
    '.c .h .cpp .hpp .cs .java .go .rs .rb .php .swift .kt .r .ipynb').split())


def extract_text(name, data):
    if not 0 < len(data) <= MAX_FILE_BYTES:
        raise ValueError('Choose a nonempty file no larger than 8 MiB.')
    suffix = PurePath(name).suffix.lower()
    if suffix in {'.xlsx', '.xlsm', '.xls', '.ods', '.odt', '.pptx'}:
        return _office(suffix, data)
    if suffix == '.pdf':
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError('PDF reading is unavailable. Run SETUP_HANNI.cmd to install the reader.') from None
        try:
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise ValueError('Unlock this PDF before attaching it.')
            if len(reader.pages) > 200:
                raise ValueError('This PDF exceeds 200 pages. Split it into smaller documents.')
            chunks, used = [], 0
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ''
                chunks.append(f'[Page {number}]\n{text}')
                used += len(chunks[-1])
                if used > MAX_TEXT:
                    break
            if not any((chunk.partition('\n')[2]).strip() for chunk in chunks):
                raise ValueError('This PDF has no readable text. Run OCR first, then attach it again.')
            text = '\n\n'.join(chunks)
            return text[:MAX_TEXT], 'application/pdf', len(text) > MAX_TEXT
        except ValueError:
            raise
        except Exception:
            raise ValueError('Could not read this PDF. Check that it is a valid, unlocked document.') from None
    if suffix == '.docx':
        try:
            with _archive(data) as archive:
                raw = archive.read('word/document.xml')
                if len(raw) > 8 * 1024 * 1024 or b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                    raise ValueError('This Word document has unsupported XML content.')
                root = ET.fromstring(raw, forbid_dtd=True)
                ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
                text = '\n'.join(''.join(node.text or '' if node.tag == ns + 't'
                    else '\t' if node.tag == ns + 'tab' else '\n' if node.tag == ns + 'br' else ''
                    for node in paragraph.iter()) for paragraph in root.iter(ns + 'p'))
            if not text.strip():
                raise ValueError('This Word document has no readable text.')
            return text[:MAX_TEXT], 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', len(text) > MAX_TEXT
        except (BadZipFile, KeyError, ET.ParseError, DefusedXmlException, RuntimeError, NotImplementedError, OSError):
            raise ValueError('Could not read this Word document. Use a valid .docx file.') from None
    if suffix not in TEXT_EXTENSIONS:
        raise ValueError('Unsupported file type. Choose PDF, DOCX, PPTX, XLSX, XLSM, XLS, ODS, ODT, HTML, text, CSV, JSON, or code. Convert legacy DOC/PPT, XLSB, and RTF first.')
    try:
        encoding = 'utf-32' if data.startswith((b'\xff\xfe\x00\x00', b'\x00\x00\xfe\xff')) else 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        text = data.decode(encoding)
    except UnicodeError:
        raise ValueError('This file is not readable text. Save it as UTF-8 and try again.') from None
    if '\x00' in text or any(ord(char) < 32 and char not in '\n\r\t\f' for char in text):
        raise ValueError('This file contains binary data. Choose a text document instead.')
    if not text.strip():
        raise ValueError('This file has no readable text.')
    text = text.replace('\r\n', '\n')
    if suffix in {'.html', '.htm'}:
        parser = _HTMLText()
        parser.feed(text)
        text = ''.join(parser.parts).strip()
        if not text:
            raise ValueError('This HTML file has no readable text. Export a static page with its content included.')
    media = {'.html': 'text/html', '.htm': 'text/html', '.csv': 'text/csv',
             '.tsv': 'text/tab-separated-values', '.json': 'application/json',
             '.jsonl': 'application/x-ndjson'}.get(suffix, 'text/plain')
    return text[:MAX_TEXT], media, len(text) > MAX_TEXT


def _archive(data):
    """Bound inflation before any XML parser or workbook loader sees the archive."""
    archive = ZipFile(BytesIO(data))
    items = archive.infolist()
    if (len(items) > 4096 or sum(info.file_size for info in items) > 24 * 1024 * 1024
            or any(info.file_size > 8 * 1024 * 1024 or info.flag_bits & 1 for info in items)):
        archive.close()
        raise ValueError('Document exceeds the 24 MiB expanded/8 MiB part reading limit or is encrypted. Split or unlock it first.')
    return archive


def _xml(archive, name):
    return ET.fromstring(archive.read(name), forbid_dtd=True)


class _HTMLText(HTMLParser):
    """Readable static text only: no rendering, scripts, or network requests."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.skipped = [], []

    def handle_starttag(self, tag, attrs):
        if tag in {'script', 'style', 'template', 'svg'}:
            self.skipped.append(tag)
        if self.skipped:
            return
        if tag in {'p', 'div', 'br', 'tr', 'li', 'h1', 'h2', 'h3', 'h4', 'section', 'article', 'title'}:
            self.parts.append('\n')
        if tag in {'td', 'th'}:
            self.parts.append('\t')

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if self.skipped:
            if tag == self.skipped[-1]:
                self.skipped.pop()
            return
        if tag in {'p', 'div', 'tr', 'li', 'h1', 'h2', 'h3', 'h4', 'section', 'article', 'title'}:
            self.parts.append('\n')

    def handle_data(self, data):
        if not self.skipped:
            self.parts.append(data)


class _Output:
    def __init__(self):
        self.parts, self.used, self.truncated, self.readable = [], 0, False, False

    def add(self, text, *, content=False):
        self.readable |= content and bool(text.strip())
        remaining = MAX_TEXT - self.used
        self.parts.append(text[:remaining])
        self.used += min(remaining, len(text))
        self.truncated |= len(text) > remaining
        return self.used < MAX_TEXT


def _cell(value):
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if hasattr(value, 'text') and hasattr(value, 'ref'):
        return {'expression': value.text, 'range': value.ref}
    return str(value)


def _excel(data, out, dimensions):
    from contextlib import ExitStack
    from openpyxl import load_workbook
    # Formula and cached-value views are both read-only. Neither executes formulas,
    # macros, external links, or calculation engines.
    with ExitStack() as stack:
        books = []
        for cached in (False, True):
            book = load_workbook(BytesIO(data), read_only=True, data_only=cached, keep_links=False)
            stack.callback(book.close)
            books.append(book)
        out.add('Cells include formulas and last-saved cached values when available. No recalculation; caches can be stale. Formatting/charts/images are not rendered.\n')
        for sheet, cache in zip(books[0].worksheets, books[1].worksheets):
            out.add('\n[Sheet ' + json.dumps(sheet.title, ensure_ascii=False) + f'; {sheet.sheet_state}]\n')
            # Ignore incorrect producer dimensions. Count real populated cells,
            # while bounding row/column expansion of sparse or hostile sheets.
            sheet.reset_dimensions(); cache.reset_dimensions()
            # openpyxl exposes the archive path on read-only sheets. Dimensions
            # were measured from actual cell references in the XML preflight.
            rows, columns = dimensions.get(sheet._worksheet_path, (0, 0))
            out.truncated |= rows > 10000 or columns > 256
            if not rows or not columns:
                continue
            bounds = dict(max_row=min(rows, 10000), max_col=min(columns, 256))
            for row, cached_row in zip(sheet.iter_rows(**bounds), cache.iter_rows(**bounds)):
                cells = {}
                for cell, saved in zip(row, cached_row):
                    if cell.value is not None:
                        value = _cell(cell.value)
                        if cell.data_type == 'f':
                            value = {'formula': value, 'cached_value': _cell(saved.value)}
                        cells[cell.coordinate] = value
                if cells and not out.add(json.dumps(cells, ensure_ascii=False) + '\n', content=True):
                    out.truncated = True
                    return


def _xls(data, out):
    import xlrd
    book = xlrd.open_workbook(file_contents=data, on_demand=True, ragged_rows=True)
    try:
        out.add('Legacy XLS: last-saved values only; formula expressions are unavailable and results may be stale. No recalculation.\n')
        for sheet in book.sheets():
            out.add('\n[Sheet ' + json.dumps(sheet.name, ensure_ascii=False) + ']\n')
            out.truncated |= sheet.nrows > 10000 or sheet.ncols > 256
            for r in range(min(sheet.nrows, 10000)):
                cells = {}
                for c in range(min(sheet.row_len(r), 256)):
                    cell = sheet.cell(r, c)
                    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, '#ERROR')
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    cells[xlrd.formula.colname(c) + str(r + 1)] = value
                if cells and not out.add(json.dumps(cells, ensure_ascii=False) + '\n', content=True):
                    out.truncated = True
                    return
    finally:
        book.release_resources()


_TEXT_NS = '{urn:oasis:names:tc:opendocument:xmlns:text:1.0}'
_TABLE_NS = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
_OFFICE_NS = '{urn:oasis:names:tc:opendocument:xmlns:office:1.0}'


def _od_text(node):
    return ''.join(_od_fragments(node))


def _od_fragments(node):
    if node.tag == _TEXT_NS + 's':
        yield ' ' * min(1000, int(node.get(_TEXT_NS + 'c', '1')))
    elif node.tag == _TEXT_NS + 'tab':
        yield '\t'
    elif node.tag == _TEXT_NS + 'line-break':
        yield '\n'
    else:
        yield node.text or ''
        for child in node:
            yield from _od_fragments(child)
            yield child.tail or ''


def _ods(root, out):
    out.add('ODS cells: formulas and saved values/text, without recalculation. Repeated rows/cells use explicit ranges.\n')
    for sheet in root.iter(_TABLE_NS + 'table'):
        out.add('\n[Sheet ' + json.dumps(sheet.get(_TABLE_NS + 'name', ''), ensure_ascii=False) + ']\n')
        row_number = 1
        for row in sheet.iter(_TABLE_NS + 'table-row'):
            repeats = int(row.get(_TABLE_NS + 'number-rows-repeated', '1'))
            if repeats < 1:
                raise ValueError('Invalid repeated row count.')
            cells, column = [], 1
            for cell in row:
                if cell.tag not in {_TABLE_NS + 'table-cell', _TABLE_NS + 'covered-table-cell'}:
                    continue
                count = int(cell.get(_TABLE_NS + 'number-columns-repeated', '1'))
                if count < 1:
                    raise ValueError('Invalid repeated cell count.')
                value = '\n'.join(_od_text(p) for p in cell.iter(_TEXT_NS + 'p'))
                values = {key: cell.get(_OFFICE_NS + key) for key in
                          ['value-type', 'value', 'date-value', 'time-value', 'boolean-value', 'currency', 'string-value']
                          if cell.get(_OFFICE_NS + key) is not None}
                if value or values or cell.get(_TABLE_NS + 'formula'):
                    cells.append({'column': column, 'repeat': count, 'text': value,
                                  **values, **({'formula': cell.get(_TABLE_NS + 'formula')} if cell.get(_TABLE_NS + 'formula') else {})})
                column += count
            if cells and not out.add(json.dumps({'row': row_number, 'repeat': repeats, 'cells': cells}, ensure_ascii=False) + '\n', content=True):
                out.truncated = True
                return
            row_number += repeats


def _office(suffix, data):
    out = _Output()
    media = {'.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
             '.xlsm': 'application/vnd.ms-excel.sheet.macroEnabled.12', '.xls': 'application/vnd.ms-excel',
             '.ods': 'application/vnd.oasis.opendocument.spreadsheet', '.odt': 'application/vnd.oasis.opendocument.text',
             '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation'}[suffix]
    try:
        if suffix == '.xls':
            _xls(data, out)
        else:
            with _archive(data) as archive:
                if suffix in {'.xlsx', '.xlsm'}:
                    # Preflight all XML with hardened parser before openpyxl uses it.
                    from openpyxl.utils.cell import coordinate_to_tuple
                    dimensions = {}
                    for name in archive.namelist():
                        if name.endswith(('.xml', '.rels')):
                            root = _xml(archive, name)
                            if name.startswith('xl/worksheets/'):
                                rows, columns = 0, 0
                                for cell in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c'):
                                    if cell.get('r'):
                                        row, column = coordinate_to_tuple(cell.get('r'))
                                        rows, columns = max(rows, row), max(columns, column)
                                dimensions[name] = rows, columns
                    _excel(data, out, dimensions)
                elif suffix in {'.ods', '.odt'}:
                    root = _xml(archive, 'content.xml')
                    if suffix == '.ods':
                        _ods(root, out)
                    else:
                        for node in root.iter():
                            if node.tag in {_TEXT_NS + 'p', _TEXT_NS + 'h'}:
                                if not out.add(_od_text(node) + '\n', content=True):
                                    out.truncated = True
                                    break
                else:
                    # Use presentation relationships, not ZIP/numeric order: slides
                    # can be reordered without renaming slideN.xml.
                    import posixpath
                    root = _xml(archive, 'ppt/presentation.xml')
                    rels = {node.get('Id'): node for node in _xml(archive, 'ppt/_rels/presentation.xml.rels')}
                    ns = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
                    for number, slide in enumerate(root.iter(ns + 'sldId'), 1):
                        rel = rels[slide.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
                        if rel.get('TargetMode') == 'External':
                            raise ValueError('External slide references are unsupported.')
                        target = rel.get('Target', '')
                        name = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'ppt/' + target)
                        if not name.startswith('ppt/slides/'):
                            raise ValueError('Invalid slide reference.')
                        page = _xml(archive, name)
                        out.add(f'\n[Slide {number}]\n')
                        for paragraph in page.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}p'):
                            drawing = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
                            text = ''.join(node.text or '' if node.tag == drawing + 't'
                                           else '\n' if node.tag == drawing + 'br' else ''
                                           for node in paragraph.iter())
                            if not out.add(text + '\n', content=True):
                                break
                        if out.used >= MAX_TEXT:
                            out.truncated = True
                            break
    except ImportError:
        raise ValueError('Document reading dependencies are unavailable. Run SETUP_HANNI.cmd to install them.') from None
    except ValueError:
        raise
    except Exception:
        raise ValueError(f'Could not read this {suffix[1:].upper()} document. Check the format and unlock it before attaching.') from None
    if not out.readable:
        raise ValueError('This document has no readable text or cells. Export its text or run OCR first.')
    return ''.join(out.parts), media, out.truncated
